import os
import time
from pptx import Presentation
from PIL import Image, ImageDraw
from supabase import create_client, Client
from google import genai
from google.genai import types
from config import (
    SUPABASE_URL,
    SUPABASE_KEY,
    GEMINI_API_KEY,
    GEMINI_EMBEDDING_MODEL,
    EMBEDDING_DIM,
    SLIDES_DIR
)

supabase: Client = create_client(SUPABASE_URL, SUPABASE_KEY)
client = genai.Client(api_key=GEMINI_API_KEY)

SLIDES_FILE = "lecture_slides.pptx"
SOURCE_ID = "slides_001"

def draw_slide_card(title: str, bullets: list, slide_num: int, output_path: str):
    """Draw a clean slide preview image so the UI can pop it up when clicked."""
    width, height = 960, 540
    img = Image.new("RGB", (width, height), color=(30, 41, 59)) 
    draw = ImageDraw.Draw(img)
    draw.rectangle([0, 0, width, 6], fill=(59, 130, 246))

    draw.text((40, 40), f"Slide {slide_num}: {title}", fill=(248, 250, 252))
    draw.line([(40, 85), (width - 40, 85)], fill=(71, 85, 105), width=2)

    y = 120
    for bullet in bullets:
        words = bullet.split()
        line = "• "
        for word in words:
            if len(line) + len(word) > 75:
                draw.text((50, y), line, fill=(226, 232, 240))
                y += 28
                line = "  " + word + " "
            else:
                line += word + " "
        if line.strip():
            draw.text((50, y), line, fill=(226, 232, 240))
            y += 36

    draw.text((40, height - 35), "Optics Course Material | Lecture Slide Preview", fill=(148, 163, 184))
    img.save(output_path)

def ingest_slides():
    print("Registering slide deck in Supabase...")
    supabase.table("sources").upsert({
        "id": SOURCE_ID,
        "title": "Optics Lecture Slides",
        "modality": "slide"
    }).execute()

    prs = Presentation(SLIDES_FILE)
    print(f"Reading {SLIDES_FILE} ({len(prs.slides)} slides found)...")

    for idx, slide in enumerate(prs.slides):
        slide_num = idx + 1
        slide_title = f"Optics Topic {slide_num}"
        bullet_list = []
        text_content = ""

        for shape in slide.shapes:
            if shape.has_text_frame:
                if shape == slide.shapes.title and shape.text.strip():
                    slide_title = shape.text.strip()
                for p in shape.text_frame.paragraphs:
                    t = p.text.strip()
                    if t:
                        bullet_list.append(t)
                        text_content += t + "\n"

        if not text_content.strip():
            continue
        slide_rel = f"assets/slides/slide_{slide_num}.png"
        slide_full = os.path.join(SLIDES_DIR, f"slide_{slide_num}.png")
        draw_slide_card(slide_title, bullet_list, slide_num, slide_full)

        try:
            resp = client.models.embed_content(
                model=GEMINI_EMBEDDING_MODEL,
                contents=text_content.strip(),
                config=types.EmbedContentConfig(output_dimensionality=EMBEDDING_DIM)
            )
            embedding = resp.embeddings[0].values
            chunk_id = f"{SOURCE_ID}_s{slide_num}"

            supabase.table("chunks").upsert({
                "id": chunk_id,
                "source_id": SOURCE_ID,
                "modality": "slide",
                "text": text_content.strip(),
                "slide": slide_num,
                "image_ref": slide_rel,
                "embedding": embedding
            }).execute()

            print(f"  [{slide_num}/{len(prs.slides)}] Processed slide {slide_num}")
            time.sleep(1)

        except Exception as err:
            print(f"  [!] Error on slide {slide_num}: {err}")
            time.sleep(2)

    print("Finished ingesting slides.")

if __name__ == "__main__":
    ingest_slides()